"""Excel report with 'User Risk' and 'Findings' sheets."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
LEVEL_FILL = {
    "HIGH": PatternFill("solid", fgColor="F8CBAD"),
    "MEDIUM": PatternFill("solid", fgColor="FFE699"),
    "LOW": PatternFill("solid", fgColor="C6E0B4"),
}


def _write_sheet(ws, df, level_col=None):
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for row in df.itertuples(index=False):
        ws.append(list(row))
    if level_col is not None:
        idx = list(df.columns).index(level_col) + 1
        for r in range(2, ws.max_row + 1):
            fill = LEVEL_FILL.get(ws.cell(row=r, column=idx).value)
            if fill:
                ws.cell(row=r, column=idx).fill = fill
    for i, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col]])
        ws.column_dimensions[get_column_letter(i)].width = min(longest + 2, 70)
    ws.freeze_panes = "A2"


def write_report(result, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "User Risk"
    _write_sheet(ws, result.risk, level_col="risk_level")
    _write_sheet(wb.create_sheet("Findings"), result.findings, level_col="severity")
    wb.save(path)
    return path
