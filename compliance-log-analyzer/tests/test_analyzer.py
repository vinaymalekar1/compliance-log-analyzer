from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from analyzer.api import create_app
from analyzer.engine import REQUIRED_COLUMNS, analyze, normalize
from analyzer.rules import Thresholds

BASE = datetime(2026, 3, 2, 10, 0)  # a Monday, 10:00
SAMPLE = Path(__file__).resolve().parent.parent / "sample_data" / "audit_log.csv"


def make_df(rows):
    """rows: (minutes_from_base, user, action, status, ip, record_count)"""
    data = [{
        "timestamp": BASE + timedelta(minutes=m), "user": u, "action": a,
        "ip_address": ip, "status": s, "record_count": rc,
    } for m, u, a, s, ip, rc in rows]
    return normalize(pd.DataFrame(data))


def rules_for(result, user):
    f = result.findings
    return list(f[f["user"] == user]["rule"])


# ---------- detection rules and edge cases ----------

def test_missing_columns_raises_value_error():
    with pytest.raises(ValueError, match="Missing required column"):
        normalize(pd.DataFrame({"timestamp": [BASE], "user": ["a"]}))


def test_failed_login_burst_detected():
    rows = [(i, "alice", "LOGIN", "FAILURE", "1.1.1.1", 0) for i in range(5)]
    result = analyze(make_df(rows))
    assert rules_for(result, "alice") == ["failed_login_burst"]


def test_failed_logins_below_threshold_or_spread_out_not_flagged():
    four = [(i, "bob", "LOGIN", "FAILURE", "1.1.1.1", 0) for i in range(4)]
    spread = [(i * 30, "carl", "LOGIN", "FAILURE", "1.1.1.1", 0) for i in range(6)]
    result = analyze(make_df(four + spread))
    assert result.findings.empty


def test_privilege_change_detected():
    result = analyze(make_df([(0, "alice", "PRIVILEGE_CHANGE", "SUCCESS", "1.1.1.1", 0)]))
    assert rules_for(result, "alice") == ["privilege_change"]


def test_bulk_export_detected_only_above_threshold():
    rows = [
        (0, "bob", "EXPORT", "SUCCESS", "1.1.1.1", 5000),
        (5, "carl", "EXPORT", "SUCCESS", "1.1.1.1", 50),
    ]
    result = analyze(make_df(rows))
    assert rules_for(result, "bob") == ["bulk_export"]
    assert rules_for(result, "carl") == []


def test_multi_ip_access_detected():
    rows = [(i * 10, "carol", "VIEW", "SUCCESS", ip, 0)
            for i, ip in enumerate(["1.1.1.1", "2.2.2.2", "3.3.3.3"])]
    result = analyze(make_df(rows))
    assert rules_for(result, "carol") == ["multi_ip_access"]


def test_off_hours_activity_detected():
    # BASE is 10:00, so +13h = 23:00 (off hours) and +1h = 11:00 (normal)
    rows = [
        (13 * 60, "eve", "VIEW", "SUCCESS", "1.1.1.1", 0),
        (60, "dave", "VIEW", "SUCCESS", "1.1.1.1", 0),
    ]
    result = analyze(make_df(rows))
    assert rules_for(result, "eve") == ["off_hours_activity"]
    assert rules_for(result, "dave") == []


def test_risk_score_is_capped_at_100_and_levels_assigned():
    rows = [(i, "alice", "PRIVILEGE_CHANGE", "SUCCESS", "1.1.1.1", 0) for i in range(6)]
    rows.append((13 * 60, "bob", "VIEW", "SUCCESS", "1.1.1.1", 0))  # only off-hours = 15 pts
    risk = analyze(make_df(rows)).risk.set_index("user")
    assert risk.loc["alice", "risk_score"] == 100 and risk.loc["alice", "risk_level"] == "HIGH"
    assert risk.loc["bob", "risk_score"] == 15 and risk.loc["bob", "risk_level"] == "LOW"
    assert list(analyze(make_df(rows)).risk["user"])[0] == "alice"  # ranked highest first


def test_empty_log_produces_no_findings():
    result = analyze(normalize(pd.DataFrame(columns=REQUIRED_COLUMNS)), Thresholds())
    assert result.findings.empty and result.risk.empty and result.total_events == 0


# ---------- API behaviour ----------

def test_api_upload_findings_risk_and_report(tmp_path):
    client = TestClient(create_app(tmp_path / "t.db", tmp_path / "reports", api_key=""))

    with open(SAMPLE, "rb") as f:
        resp = client.post("/upload", files={"file": ("audit_log.csv", f, "text/csv")})
    assert resp.status_code == 200
    run_id = resp.json()["run_id"]
    assert resp.json()["total_findings"] > 0

    findings = client.get(f"/runs/{run_id}/findings").json()
    assert {x["rule"] for x in findings} >= {"failed_login_burst", "privilege_change"}
    assert all(x["user"] == "alice" for x in client.get(f"/runs/{run_id}/findings?user=alice").json())

    scores = [r["risk_score"] for r in client.get(f"/runs/{run_id}/risk").json()]
    assert scores == sorted(scores, reverse=True)

    bad = client.post("/upload", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert bad.status_code == 400
    assert client.get("/runs/999/findings").status_code == 404

    report = client.get(f"/runs/{run_id}/report")
    assert report.status_code == 200
    out = tmp_path / "downloaded.xlsx"
    out.write_bytes(report.content)
    assert load_workbook(out).sheetnames == ["User Risk", "Findings"]


def test_api_key_required_when_configured(tmp_path):
    client = TestClient(create_app(tmp_path / "t.db", tmp_path / "reports", api_key="secret123"))
    assert client.get("/health").status_code == 200  # health stays open
    assert client.get("/runs").status_code == 401
    assert client.get("/runs", headers={"X-API-Key": "wrong"}).status_code == 401
    assert client.get("/runs", headers={"X-API-Key": "secret123"}).status_code == 200
